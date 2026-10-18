import numpy as np
import openpyxl
import pytest

from easycellcounting.io import natural_key
from easycellcounting.report import animal_id, region_code, results_table, section_label, write_excel
from easycellcounting.roi import ROI, clip_box, crop, load_rois, rotate_image, save_rois, scale_box


def test_scale_box_uses_width_for_x_and_height_for_y():
    # preview 100 x 200 (h x w) of a 1000 x 4000 image: x scales by 20, y by 10
    assert scale_box((10, 5, 20, 10), (100, 200), (1000, 4000)) == (200, 50, 400, 100)


def test_clip_box():
    assert clip_box((-5, -5, 20, 20), (10, 10)) == (0, 0, 10, 10)
    with pytest.raises(ValueError):
        clip_box((50, 50, 5, 5), (10, 10))


def test_rotation_and_crop():
    img = np.zeros((40, 60, 3), np.uint8)
    img[5:10, 5:10] = 200
    assert rotate_image(img, 360) is img
    rot = rotate_image(img, 180)
    ys, xs = np.nonzero(rot[:, :, 0] > 100)               # the square moved to the far corner
    assert 29 <= ys.mean() <= 34 and 49 <= xs.mean() <= 54
    roi = ROI("a.tif", 48, 28, 9, 9, rotation=180)
    assert crop(img, roi).sum() == rot.sum()


def test_roi_file_roundtrip(tmp_path):
    rois = [ROI("a 1.tif", 1, 2, 30, 40, 7.0, 200, 40, 900, 1e6, "custom"),
            ROI("b.tif", 0, 0, 10, 10)]
    p = save_rois(rois, tmp_path / "rois.csv")
    assert load_rois(p) == rois
    assert rois[0].params().threshold == 200
    assert rois[0].params().cluster_max_area == 1e6


def test_natural_sort():
    names = ["x_s10.tif", "x_s2.tif", "x_s1.tif"]
    assert sorted(names, key=natural_key) == ["x_s1.tif", "x_s2.tif", "x_s10.tif"]


def test_file_name_parsing():
    name = "Zymo6 3730.4 1-12_s1.tif"
    assert animal_id(name) == "3730.4"
    assert animal_id(name, slice_="6:13") == "3730.4"           # the original's slice
    assert animal_id("CFA 30.14 01-12_s1.tif") == "30.14"
    assert animal_id("mouseA_s3.tif", pattern=None) == "mouseA"
    assert animal_id("M12 slice_s3.tif", pattern=r"M(\d+)") == "12"
    assert section_label(name) == "s1"
    assert region_code("Zymo6 3730.4 B2_s1.tif") == "B2"
    assert region_code("Zymo6 3730.4 CA1_s1.tif") == ""          # no false match inside "CA1"


def test_excel_layout(tmp_path):
    rows = [{"file": f"Z {a} {r}_s{s}.tif", "left": s, "right": 10 * s, "total": 11 * s}
            for a, r in (("3730.4", "A1"), ("3731.2", "C4")) for s in (1, 2, 3)]
    df = results_table(rows)
    assert list(df["section"]) == [1, 2, 3, 1, 2, 3]
    p = write_excel(df, tmp_path / "counts.xlsx", {"threshold": 160})
    wb = openpyxl.load_workbook(p)
    ws = wb["Left-Right counts"]
    assert ws["A1"].value == "Left counts"
    assert [c.value for c in ws[3]][:4] == ["3730.4", 1, 2, 3]
    assert ws["A6"].value == "Right counts"
    assert [c.value for c in ws[9]][:4] == ["3731.2", 10, 20, 30]
    assert ws["B3"].fill.fgColor.indexed == 27                   # A1 colour
    assert set(wb.sheetnames) == {"Left-Right counts", "Per image", "Parameters"}
