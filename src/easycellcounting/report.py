"""
Results tables: a tidy per-image CSV and an Excel workbook laid out like the original
(one row per animal, sections across, left and right tables, colour-coded regions).
"""
from __future__ import annotations

from pathlib import Path
import re

import pandas as pd

# Excel indexed colours used by the original script for the region codes in file names
REGION_COLOURS = {"A1": 27, "A2": 44, "A3": 49, "A4": 4,
                  "B1": 5, "B2": 50, "B3": 57, "B4": 19,
                  "C1": 45, "C2": 29, "C3": 22, "C4": 23}
REGION_PATTERN = re.compile(r"(?<![A-Za-z0-9])([A-C][1-4])(?![0-9])")
DEFAULT_ANIMAL_PATTERN = r"\d+\.\d+"


def animal_id(name: str, pattern: str | None = DEFAULT_ANIMAL_PATTERN,
              slice_: str | None = None) -> str:
    """Animal tag of an image file name.

    `pattern`: regex, first match (or its first group) is the animal, e.g. the
    default ``\\d+\\.\\d+`` gives "3730.4" for "Zymo6 3730.4 1-12_s1.tif".
    `slice_`: "start:stop" character slice instead (the original used 6:13)."""
    if slice_:
        a, b = (int(v) if v else None for v in slice_.split(":"))
        return name[a:b].strip()
    if pattern:
        m = re.search(pattern, name)
        if m:
            return (m.group(1) if m.groups() else m.group(0)).strip()
    return re.sub(r"[_ -]?s\d+$", "", Path(name).stem)


def section_label(name: str) -> str:
    """Section tag at the end of a file name ("s12" in "... 1-12_s12.tif"), or ""."""
    m = re.search(r"[_ -](s\d+)$", Path(name).stem)
    return m.group(1) if m else ""


def region_code(name: str) -> str:
    m = REGION_PATTERN.search(Path(name).stem)
    return m.group(1) if m else ""


def results_table(rows, animal_pattern=DEFAULT_ANIMAL_PATTERN, animal_slice=None) -> pd.DataFrame:
    """Tidy table, one row per image: animal, section (order of the image within its
    animal, as in the original workbook), the section tag and region code from the name."""
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    df.insert(1, "animal", [animal_id(f, animal_pattern, animal_slice) for f in df["file"]])
    df.insert(2, "section", df.groupby("animal", sort=False).cumcount() + 1)
    df.insert(3, "section_label", [section_label(f) for f in df["file"]])
    df.insert(4, "region", [region_code(f) for f in df["file"]])
    return df


def write_csv(df: pd.DataFrame, path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)
    return path


def write_excel(df: pd.DataFrame, path, parameters: dict | None = None) -> Path:
    """Workbook with the left/right summary, the per-image table and the run parameters."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.styles.colors import Color

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Left-Right counts"
    bold = Font(bold=True)
    animals = list(dict.fromkeys(df["animal"])) if len(df) else []
    n_sections = int(df["section"].max()) if len(df) else 0

    def table(top, title, column):
        ws.cell(row=top, column=1, value=title).font = bold
        ws.cell(row=top + 1, column=1, value="Animal").font = bold
        for s in range(1, n_sections + 1):
            ws.cell(row=top + 1, column=s + 1, value=s).font = bold
        for i, animal in enumerate(animals):
            r = top + 2 + i
            ws.cell(row=r, column=1, value=animal)
            for _, rec in df[df["animal"] == animal].iterrows():
                cell = ws.cell(row=r, column=int(rec["section"]) + 1, value=int(rec[column]))
                if rec["region"] in REGION_COLOURS:
                    cell.fill = PatternFill(patternType="solid",
                                            fgColor=Color(indexed=REGION_COLOURS[rec["region"]]))
        return top + 2 + len(animals)

    end = table(1, "Left counts", "left")
    end = table(end + 1, "Right counts", "right")
    end = table(end + 1, "Total counts", "total")
    legend = end + 1
    ws.cell(row=legend, column=1, value="Region colour code").font = bold
    for j, (code, colour) in enumerate(REGION_COLOURS.items()):
        c = ws.cell(row=legend, column=j + 2, value=code)
        c.fill = PatternFill(patternType="solid", fgColor=Color(indexed=colour))
    ws.column_dimensions["A"].width = 16

    tidy = wb.create_sheet("Per image")
    tidy.append(list(df.columns))
    for c in tidy[1]:
        c.font = bold
    for rec in df.itertuples(index=False):
        tidy.append([None if (isinstance(v, float) and v != v) else v for v in rec])

    if parameters:
        ps = wb.create_sheet("Parameters")
        for k, v in parameters.items():
            ps.append([k, str(v)])
    wb.save(path)
    return path
